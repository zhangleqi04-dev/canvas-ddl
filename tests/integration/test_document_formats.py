from datetime import datetime
from io import BytesIO
import hashlib
import json
from types import SimpleNamespace

import pytest

from canvas_ddl.courses.models import Course
from canvas_ddl.deadlines.query import DeadlineQuery
from canvas_ddl.deadlines.service import DeadlineService
from canvas_ddl.deadlines.time_intent import TimeIntent
from canvas_ddl.documents.ingestion import DocumentIngestionService
from canvas_ddl.documents.ocr import OcrPageResult
from canvas_ddl.documents.parser import DocumentParser
from canvas_ddl.documents.registry import OfficialDocumentRegistry
from canvas_ddl.documents.repository import DocumentRepository
from canvas_ddl.documents.preparation import DocumentPreparationService
from canvas_ddl.documents.refresh import DocumentLibraryRefresher
from canvas_ddl.documents.inventory import CourseDocumentInventory
from canvas_ddl.serialization import encode


NOW = datetime.fromisoformat("2026-09-17T13:05:00+08:00")
COURSE = Course("12345", "CS3244", "Machine Learning")
LINE = "Midterm Exam: 2026-09-23 14:00"


def docx_bytes():
    from docx import Document
    stream = BytesIO()
    document = Document()
    document.add_paragraph(LINE)
    document.save(stream)
    return stream.getvalue()


def pptx_bytes():
    from pptx import Presentation
    from pptx.util import Inches
    stream = BytesIO()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = LINE
    deck.save(stream)
    return stream.getvalue()


def xlsx_bytes():
    from openpyxl import Workbook
    stream = BytesIO()
    workbook = Workbook()
    workbook.active.title = "Assessment"
    workbook.active.append(["Midterm Exam", "2026-09-23", "14:00"])
    workbook.save(stream)
    return stream.getvalue()


FORMATS = {
    "outline.docx": docx_bytes,
    "slides.pptx": pptx_bytes,
    "schedule.xlsx": xlsx_bytes,
    "schedule.csv": lambda: (LINE + "\n").encode(),
    "schedule.txt": lambda: (LINE + "\n").encode(),
    "schedule.md": lambda: (LINE + "\n").encode(),
    "schedule.rtf": lambda: (r"{\rtf1\ansi " + LINE + "}").encode(),
    "syllabus.html": lambda: ("<html><body><table><tr><td>Midterm Exam</td><td>2026-09-23 14:00</td></tr>"
                                "<script>ignore()</script></table></body></html>").encode(),
}


class NoRows:
    source_type = "canvas_assignment"
    def collect(self, **_kwargs):
        return []


class Client:
    def get(self, *_args, **_kwargs):
        raise AssertionError("No Canvas detail request expected")
    def close(self):
        pass


def setup_document(tmp_path, name, content, *, parser=None):
    path = tmp_path / name
    path.write_bytes(content)
    digest = hashlib.sha256(content).hexdigest()
    registry_path = tmp_path / "registry.json"
    registry_path.write_text(json.dumps({"documents": [{
        "document_id": "outline", "course_id": COURSE.course_id, "course_code": COURSE.course_code,
        "course_name": COURSE.course_name, "document_name": name,
        "document_kind": "official_course_document",
        "source_url": "https://canvas.example/courses/12345/files/99", "path": name,
        "sha256": digest, "approved_by": "fixture-human",
        "valid_from": "2026-08-01", "valid_until": "2026-12-31"
    }]}), encoding="utf8")
    registry = OfficialDocumentRegistry(registry_path, canvas_origin="https://canvas.example")
    ingestion = DocumentIngestionService(registry, DocumentRepository(tmp_path / "documents.sqlite3"),
        timezone="Asia/Singapore", base_url="https://canvas.example", clock=lambda: NOW, parser=parser)
    return path, ingestion


@pytest.mark.parametrize("name,factory", FORMATS.items())
def test_supported_office_and_text_formats_enter_the_same_validated_pipeline(tmp_path, name, factory):
    _, ingestion = setup_document(tmp_path, name, factory())
    report = ingestion.ingest("outline", COURSE)
    assert report["confirmed"] == 1
    deadlines, summary, warnings, complete = ingestion.load_deadlines([COURSE])
    assert complete and not warnings and len(deadlines) == 1
    source = deadlines[0].sources[0]
    assert source.document_name == name and source.location
    assert source.evidence_text and source.validation.status == "confirmed"


def test_docx_is_not_reopened_during_final_query(tmp_path, monkeypatch):
    path, ingestion = setup_document(tmp_path, "outline.docx", docx_bytes())
    ingestion.ingest("outline", COURSE)
    path.unlink()
    monkeypatch.setattr(ingestion.parser, "parse", lambda *_a, **_k: pytest.fail("query reopened source document"))
    service = DeadlineService(Client(), base_url="https://canvas.example",
        repository=SimpleNamespace(list_courses=lambda: [COURSE]), collectors=[NoRows()],
        document_ingestion=ingestion, clock=lambda: NOW)
    result = encode(service.query(DeadlineQuery(time_intent=TimeIntent("calendar_week", "下周", offset=1),
                                                types=("exam",), document_mode="existing")))
    assert result["count"] == 1 and result["deadlines"][0]["sources"][0]["location"] == "paragraph 1"


def test_standalone_image_uses_local_ocr_then_independent_validation(tmp_path):
    from PIL import Image
    stream = BytesIO()
    Image.new("RGB", (100, 30), "white").save(stream, format="PNG")
    class Ocr:
        def recognize_image(self, _content):
            return OcrPageResult(LINE, (0.99,))
    _, ingestion = setup_document(tmp_path, "schedule.png", stream.getvalue(),
                                  parser=DocumentParser(ocr=Ocr()))
    report = ingestion.ingest("outline", COURSE)
    assert report["confirmed"] == 1 and report["ocr_pages"] == [1]
    deadlines, _, _, _ = ingestion.load_deadlines([COURSE])
    assert deadlines[0].sources[0].ocr_engine == "PP-OCRv6_small"


def test_canvas_syllabus_and_pages_are_discovered_scanned_and_left_pending_review(tmp_path):
    class ContentClient:
        base_url = "https://canvas.example"
        def get_paginated(self, path, **_kwargs):
            if path.endswith("/files"):
                return []
            if path.endswith("/pages"):
                return [{"url": "exam-schedule", "title": "Exam Schedule", "published": True}]
            raise AssertionError(path)
        def get(self, path, **_kwargs):
            if path == "/api/v1/courses/12345":
                return {"id": 12345, "syllabus_body": f"<p>{LINE}</p>", "updated_at": "v1"}
            if path.endswith("/pages/exam-schedule"):
                return {"url": "exam-schedule", "title": "Exam Schedule",
                        "body": f"<p>{LINE}</p>", "updated_at": "v1"}
            raise AssertionError(path)
    registry_path = tmp_path / "registry.json"
    registry_path.write_text('{"documents": []}', encoding="utf8")
    client = ContentClient()
    registry = OfficialDocumentRegistry(registry_path, canvas_origin=client.base_url)
    ingestion = DocumentIngestionService(registry, DocumentRepository(tmp_path / "documents.sqlite3"),
        timezone="Asia/Singapore", base_url=client.base_url, clock=lambda: NOW)
    refresher = DocumentLibraryRefresher(DocumentPreparationService(client, registry), ingestion, clock=lambda: NOW)
    report = refresher.refresh([COURSE])
    pending = json.loads((tmp_path / "pending-review.json").read_text(encoding="utf8"))["documents"]
    assert report["scope"] == "all_course_documents" and len(pending) == 2
    assert {row["document_kind"] for row in pending} == {"canvas_syllabus", "canvas_page"}
    assert all(action["state"] == "pending_review" for action in report["actions"])
    assert all(action["scan"]["candidate_count"] == 1 for action in report["actions"])
    assert registry.list_documents() == ()
    syllabus_id = next(row["document_id"] for row in pending if row["document_kind"] == "canvas_syllabus")
    refresher.preparation.approve(syllabus_id, approved_by="fixture-human",
                                  valid_from="2026-08-01", valid_until="2026-12-31")
    assert ingestion.ingest(syllabus_id, COURSE)["confirmed"] == 1


def test_course_document_inventory_includes_supported_modern_formats_only():
    rows = [
        {"id": 1, "filename": "outline.docx"}, {"id": 2, "filename": "slides.pptx"},
        {"id": 3, "filename": "schedule.xlsx"}, {"id": 4, "filename": "dates.csv"},
        {"id": 5, "filename": "notice.png"}, {"id": 6, "filename": "notes.txt"},
        {"id": 7, "filename": "video.mp4"}, {"id": 8, "filename": "legacy.doc"},
        {"id": 9, "filename": "package.zip"},
    ]
    client = SimpleNamespace(get_paginated=lambda *_args, **_kwargs: rows)
    files, warnings = CourseDocumentInventory(client).list("12345")
    assert [row["id"] for row in files] == [1, 2, 3, 4, 5, 6]
    assert not warnings
