"""Bounded multi-format parser for approved and provisional course documents."""
from datetime import datetime
import csv
import hashlib
import logging
import zipfile
from io import BytesIO, StringIO

from canvas_ddl.canvas.errors import ApplicationError
from .formats import document_format
from .models import OfficialDocument, DocumentDraft, ParsedDocument, DocumentPage
from .ocr import OcrUnavailable, OcrFailed


class DocumentParser:
    MAX_BYTES = 30 * 1024 * 1024
    MAX_UNITS = 5000
    MAX_UNIT_TEXT = 200000
    MAX_TOTAL_TEXT = 5_000_000

    def __init__(self, *, ocr=None, native_text_threshold=16):
        self.ocr = ocr
        self.native_text_threshold = native_text_threshold

    def parse(self, document: OfficialDocument | DocumentDraft, *, now: datetime) -> ParsedDocument:
        try:
            if not 0 < document.path.stat().st_size <= self.MAX_BYTES:
                raise ValueError
            content = document.path.read_bytes()
            if hashlib.sha256(content).hexdigest() != document.sha256:
                raise ApplicationError("DOCUMENT_HASH_MISMATCH", "The document does not match its approved content hash.")
            format_name = document_format(document.document_name) or document_format(document.path.name)
            if format_name is None:
                raise ApplicationError("DOCUMENT_FORMAT_UNSUPPORTED", "The registered course document format is not supported.")
            parser = getattr(self, f"_parse_{format_name}", None)
            if parser is None:
                raise ApplicationError("DOCUMENT_FORMAT_UNSUPPORTED", "The registered course document format is not supported.")
            pages = tuple(parser(content))
            if not pages or len(pages) > self.MAX_UNITS:
                raise ValueError
            total = 0
            for page in pages:
                if len(page.text) > self.MAX_UNIT_TEXT:
                    raise ValueError
                total += len(page.text)
                if total > self.MAX_TOTAL_TEXT:
                    raise ValueError
        except ApplicationError:
            raise
        except Exception:
            raise ApplicationError("DOCUMENT_PARSE_FAILED", "The registered course document could not be parsed safely.") from None
        return ParsedDocument(document.document_id, document.sha256, pages, now)

    @staticmethod
    def _decode(content):
        return content.decode("utf-8-sig")

    @staticmethod
    def _safe_ooxml(content):
        with zipfile.ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 10000 or sum(item.file_size for item in infos) > 100 * 1024 * 1024:
                raise ValueError

    def _parse_pdf(self, content):
        from pypdf import PdfReader
        logger = logging.getLogger("pypdf")
        prior_handlers, prior_propagate = logger.handlers, logger.propagate
        logger.handlers, logger.propagate = [logging.NullHandler()], False
        try:
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError
            if not 0 < len(reader.pages) <= 300:
                raise ValueError
            pages = []
            for number, page in enumerate(reader.pages, 1):
                if page.get_contents() and len(page.get_contents().get_data()) > 16 * 1024 * 1024:
                    raise ValueError
                mode = "layout"
                try:
                    text = page.extract_text(extraction_mode=mode, layout_mode_strip_rotated=False) or ""
                except ZeroDivisionError:
                    mode = "plain"
                    text = page.extract_text(extraction_mode=mode) or ""
                confidences, engine = (), None
                if self.ocr is not None and len(text.strip()) < self.native_text_threshold:
                    try:
                        recognized = self.ocr.recognize_pdf_page(content, number - 1)
                        text, confidences, engine = recognized.text, recognized.line_confidences, recognized.engine
                        mode = "ocr_ppocrv6_small" if text.strip() else "ocr_empty"
                    except OcrUnavailable:
                        mode = "ocr_unavailable"
                    except OcrFailed:
                        mode = "ocr_failed"
                pages.append(DocumentPage(number, text, mode, confidences, engine, f"page {number}"))
            return pages
        finally:
            logger.handlers, logger.propagate = prior_handlers, prior_propagate

    def _parse_docx(self, content):
        from docx import Document
        self._safe_ooxml(content)
        document = Document(BytesIO(content))
        units, number = [], 0
        for index, paragraph in enumerate(document.paragraphs, 1):
            text = paragraph.text.strip()
            if text:
                number += 1
                units.append(DocumentPage(number, text, "docx_text", location=f"paragraph {index}"))
        for table_number, table in enumerate(document.tables, 1):
            for row_number, row in enumerate(table.rows, 1):
                text = " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells if cell.text.strip())
                if text:
                    number += 1
                    units.append(DocumentPage(number, text, "docx_table", location=f"table {table_number}, row {row_number}"))
        return units or [DocumentPage(1, "", "docx_empty", location="document")]

    def _parse_pptx(self, content):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        self._safe_ooxml(content)
        deck, units = Presentation(BytesIO(content)), []
        for slide_number, slide in enumerate(deck.slides, 1):
            lines, scores, engines, ocr_problem = [], [], [], None
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    lines.extend(line.strip() for line in shape.text.splitlines() if line.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        value = " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells if cell.text.strip())
                        if value:
                            lines.append(value)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
                lines.extend(line.strip() for line in notes.splitlines() if line.strip())
            if self.ocr is not None and len(" ".join(lines)) < self.native_text_threshold:
                native_line_count = len(lines)
                for shape in slide.shapes:
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        try:
                            recognized = self.ocr.recognize_image(shape.image.blob)
                            if recognized.text:
                                if not scores:
                                    scores = [1.0] * native_line_count
                                lines.extend(recognized.text.splitlines())
                                scores.extend(recognized.line_confidences)
                                engines.append(recognized.engine)
                        except OcrUnavailable:
                            ocr_problem = "ocr_unavailable"
                        except OcrFailed:
                            ocr_problem = "ocr_failed"
            mode = ("ocr_ppocrv6_small" if scores else ocr_problem if ocr_problem
                    else "pptx_text" if lines else "pptx_empty")
            units.append(DocumentPage(slide_number, "\n".join(lines), mode, tuple(scores),
                                      engines[0] if engines else None, f"slide {slide_number}"))
        return units

    def _parse_xlsx(self, content):
        from openpyxl import load_workbook
        self._safe_ooxml(content)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        units, number = [], 0
        try:
            if len(workbook.worksheets) > 100:
                raise ValueError
            for sheet in workbook.worksheets:
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if not values:
                        continue
                    number += 1
                    units.append(DocumentPage(number, " | ".join(values), "xlsx_row",
                                              location=f"sheet {sheet.title!r}, row {row_number}"))
                    if number > self.MAX_UNITS:
                        raise ValueError
        finally:
            workbook.close()
        return units or [DocumentPage(1, "", "xlsx_empty", location="workbook")]

    def _parse_csv(self, content):
        units = []
        for number, row in enumerate(csv.reader(StringIO(self._decode(content))), 1):
            text = " | ".join(value.strip() for value in row if value.strip())
            if text:
                units.append(DocumentPage(len(units) + 1, text, "csv_row", location=f"row {number}"))
            if len(units) > self.MAX_UNITS:
                raise ValueError
        return units or [DocumentPage(1, "", "csv_empty", location="document")]

    def _text_units(self, text, mode):
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        return [DocumentPage(index, line, mode, location=f"line {index}") for index, line in enumerate(lines, 1)] or [
            DocumentPage(1, "", mode + "_empty", location="document")]

    def _parse_text(self, content):
        return self._text_units(self._decode(content), "plain_text")

    def _parse_markdown(self, content):
        return self._text_units(self._decode(content), "markdown_text")

    def _parse_rtf(self, content):
        from striprtf.striprtf import rtf_to_text
        return self._text_units(rtf_to_text(content.decode("utf-8-sig")), "rtf_text")

    def _parse_html(self, content):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(self._decode(content), "html.parser")
        for element in soup(("script", "style", "noscript")):
            element.decompose()
        lines = []
        for row in soup.find_all("tr"):
            value = " | ".join(cell.get_text(" ", strip=True) for cell in row.find_all(("th", "td"))
                               if cell.get_text(" ", strip=True))
            if value:
                lines.append(value)
            row.decompose()
        lines.extend(element.get_text(" ", strip=True) for element in
                     soup.find_all(("h1", "h2", "h3", "h4", "h5", "h6", "p", "li"))
                     if element.get_text(" ", strip=True))
        if not lines:
            lines = [line.strip() for line in soup.get_text("\n").splitlines() if line.strip()]
        return [DocumentPage(index, line, "html_text", location=f"block {index}")
                for index, line in enumerate(lines, 1)] or [DocumentPage(1, "", "html_empty", location="document")]

    def _parse_image(self, content):
        if self.ocr is None:
            return [DocumentPage(1, "", "ocr_unavailable", location="image")]
        try:
            recognized = self.ocr.recognize_image(content)
            mode = "ocr_ppocrv6_small" if recognized.text.strip() else "ocr_empty"
            return [DocumentPage(1, recognized.text, mode, recognized.line_confidences,
                                 recognized.engine, "image")]
        except OcrUnavailable:
            return [DocumentPage(1, "", "ocr_unavailable", location="image")]
        except OcrFailed:
            return [DocumentPage(1, "", "ocr_failed", location="image")]
